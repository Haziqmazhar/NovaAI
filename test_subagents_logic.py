"""
Standalone test of the whitelist/discovery/extraction logic in subagents.py.
Like test_actions_logic.py, this has zero network/API-key dependency — it
uses fixture files in a temp directory, never the real config.json folders,
so it can't touch anything outside a throwaway temp dir.
"""
import functools
import http.server
import os
import shutil
import tempfile
import threading

import subagents
from subagents import _extract_text, _resolve_whitelisted_file, list_documents

tmp_root = tempfile.mkdtemp(prefix="nova_test_")
folder_a = os.path.join(tmp_root, "folder_a")
folder_b = os.path.join(tmp_root, "folder_b")
os.makedirs(folder_a)
os.makedirs(folder_b)

config = {"folders": {"downloads": folder_a, "documents": folder_b}}

try:
    # --- fixtures ---
    with open(os.path.join(folder_a, "notes.txt"), "w", encoding="utf-8") as f:
        f.write("hello from notes")
    with open(os.path.join(folder_a, "image.png"), "wb") as f:
        f.write(b"\x89PNG fake")  # unsupported extension, should be omitted from listings
    with open(os.path.join(folder_b, "data.csv"), "w", encoding="utf-8", newline="") as f:
        f.write("name,amount\nwidget,10\ngadget,20\n")

    from docx import Document
    doc = Document()
    doc.add_paragraph("hello from docx")
    doc.save(os.path.join(folder_b, "report.docx"))

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["col1", "col2"])
    ws.append(["x", 42])
    wb.save(os.path.join(folder_b, "sheet.xlsx"))

    print("=== _resolve_whitelisted_file ===")
    found = _resolve_whitelisted_file(config, "notes.txt")
    assert found == os.path.realpath(os.path.join(folder_a, "notes.txt")), found
    print("exact match: OK")

    # basename stripping means '..'/absolute tricks can't escape the folder —
    # they just fail to resolve rather than reading something unintended.
    outside = os.path.join(tmp_root, "secret.txt")
    with open(outside, "w") as f:
        f.write("should never be reachable")
    assert _resolve_whitelisted_file(config, "../secret.txt") is None
    assert _resolve_whitelisted_file(config, os.path.join(tmp_root, "secret.txt")) is None
    print("traversal guard: OK")

    assert _resolve_whitelisted_file(config, "nope.txt") is None
    print("missing file: OK")

    print("\n=== list_documents ===")
    single = list_documents(config, "downloads")
    print(single)
    assert single.startswith("downloads:")
    assert "notes.txt" in single
    assert "image.png" not in single  # unsupported extension omitted
    assert "unsupported formats omitted" in single

    substring = list_documents(config, "down")
    assert substring == single
    print("substring folder match: OK")

    unknown = list_documents(config, "nonexistent_folder")
    assert "isn't one of the whitelisted folders" in unknown
    print("unknown folder: OK")

    everything = list_documents(config, None)
    print(everything)
    assert "downloads:" in everything and "documents:" in everything
    assert "data.csv" in everything and "report.docx" in everything and "sheet.xlsx" in everything
    print("list-all: OK")

    print("\n=== _extract_text ===")
    csv_text = _extract_text(os.path.join(folder_b, "data.csv"))
    assert "widget" in csv_text and "10" in csv_text
    print("csv: OK")

    docx_text = _extract_text(os.path.join(folder_b, "report.docx"))
    assert "hello from docx" in docx_text
    print("docx: OK")

    xlsx_text = _extract_text(os.path.join(folder_b, "sheet.xlsx"))
    assert "Sheet1" in xlsx_text and "42" in xlsx_text
    print("xlsx: OK")

    truncated = _extract_text(os.path.join(folder_a, "notes.txt"), max_chars=5)
    assert truncated == "hello\n[...truncated...]"
    print("truncation budget respected: OK")

    print("\n=== download_file ===")
    serve_dir = os.path.join(tmp_root, "serve")
    os.makedirs(serve_dir)
    with open(os.path.join(serve_dir, "hello.txt"), "w", encoding="utf-8") as f:
        f.write("hello from the network")
    with open(os.path.join(serve_dir, "big.txt"), "w", encoding="utf-8") as f:
        f.write("x" * 100)

    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=serve_dir)
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    server_thread = threading.Thread(target=server.serve_forever, daemon=True)
    server_thread.start()
    base_url = f"http://127.0.0.1:{server.server_port}"

    try:
        result = subagents.download_file(config, f"{base_url}/hello.txt", "downloads")
        print(result)
        assert result.startswith("Downloaded 'hello.txt'"), result
        saved_path = os.path.join(folder_a, "hello.txt")
        with open(saved_path, "r", encoding="utf-8") as f:
            assert f.read() == "hello from the network"
        print("basic download: OK")

        # same URL again must not overwrite -> "hello (1).txt"
        result2 = subagents.download_file(config, f"{base_url}/hello.txt", "downloads")
        assert "hello (1).txt" in result2, result2
        assert os.path.isfile(os.path.join(folder_a, "hello (1).txt"))
        print("no-overwrite collision handling: OK")

        blocked = subagents.download_file(config, f"{base_url}/hello.txt", "downloads", file_name="evil.exe")
        assert "Refusing to download" in blocked, blocked
        assert not os.path.isfile(os.path.join(folder_a, "evil.exe"))
        print("blocked extension: OK")

        non_http = subagents.download_file(config, "ftp://example.com/x.txt", "downloads")
        assert "only http/https URLs" in non_http, non_http
        print("non-http scheme rejected: OK")

        bad_folder = subagents.download_file(config, f"{base_url}/hello.txt", "nonexistent_folder")
        assert "isn't one of the whitelisted folders" in bad_folder, bad_folder
        print("unknown folder rejected: OK")

        original_cap = subagents.MAX_DOWNLOAD_BYTES
        subagents.MAX_DOWNLOAD_BYTES = 10
        try:
            too_big = subagents.download_file(config, f"{base_url}/big.txt", "downloads")
            assert "over the" in too_big, too_big
            assert not os.path.isfile(os.path.join(folder_a, "big.txt"))
            assert not os.path.isfile(os.path.join(folder_a, "big.txt.part"))
        finally:
            subagents.MAX_DOWNLOAD_BYTES = original_cap
        print("size cap enforced: OK")
    finally:
        server.shutdown()
        server_thread.join(timeout=5)

    print("\nAll sanity assertions passed.")
finally:
    shutil.rmtree(tmp_root, ignore_errors=True)
